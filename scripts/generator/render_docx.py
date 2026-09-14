"""DOCX rendering via python-docx, with forced reproducibility."""

from __future__ import annotations

from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH

from . import reproducibility
from .models import Block

_ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
}


def render_docx(blocks: list[Block], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = docx.Document()
    for b in blocks:
        if b.kind == "heading":
            p = document.add_heading(b.text, level=2)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif b.kind == "paragraph":
            for line in b.text.split("\n"):
                p = document.add_paragraph()
                run = p.add_run(line)
                run.bold = b.bold
                if b.size == "small":
                    run.font.size = docx.shared.Pt(7)
                p.alignment = _ALIGN.get(b.align, WD_ALIGN_PARAGRAPH.LEFT)
        elif b.kind == "table":
            rows = b.rows
            t = document.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Table Grid"
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    t.cell(r, c).text = cell
        elif b.kind == "spacer":
            document.add_paragraph()
        elif b.kind == "page_break":
            document.add_page_break()
        elif b.kind == "raw_html":
            document.add_paragraph(b.raw_html)
    document.save(str(path))
    reproducibility.normalize_docx(path)
